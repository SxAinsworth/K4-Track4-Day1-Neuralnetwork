"""results_table.py — lưu kết quả từng lần chạy ra JSON và điền experiments.xlsx.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    out = Path(results_dir)
    out.mkdir(parents=True, exist_ok=True)
    path = out / f"{result['cfg']['exp_id']}.json"
    payload = {k: result[k] for k in ("cfg", "history", "summary")}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=1, allow_nan=True), encoding="utf-8")
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    files = sorted(Path(results_dir).glob("*.json"))
    return [json.loads(f.read_text(encoding="utf-8")) for f in files]


COLUMNS = ["exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay", "batch", "epochs", "hidden",
           "dropout", "clip_norm", "precision", "init", "seed", "step0_loss", "best_val_loss", "best_epoch",
           "final_train_loss", "final_val_loss", "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB",
           "diverged", "eval_acc", "eval_macro_f1", "figure_file", "notes"]
FORMULA_COLUMNS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
OPTIMIZER_NAMES = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}
MAX_ROWS = 60          # mẫu có sẵn công thức cho dòng 2..61


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng.

    eval_scores: dict có "accuracy" và "macro_f1" (đúng khoá của eval_result.json do scripts/evaluate.py ghi).
    Giá trị được chuyển sang đúng danh sách chọn của mẫu (CE/MSE, SGD+momentum, 'none', Y/N, ...).
    notes gộp: notes trong cfg + lịch lr (nếu có) + số bước GradScaler bỏ qua + notes truyền vào.
    """
    c, s = result["cfg"], result["summary"]
    note_parts = [c.get("notes") or ""]
    if c.get("scheduler"):
        note_parts.append(f"scheduler={c['scheduler']}")
    if s.get("skipped_steps"):
        note_parts.append(f"GradScaler bỏ qua {s['skipped_steps']} bước (gradient inf/NaN)")
    note_parts.append(notes)
    return {
        "exp_id": c["exp_id"], "group": c["group"], "description": c["description"],
        "loss": c["loss"].upper(), "optimizer": OPTIMIZER_NAMES[c["optimizer"]], "lr": c["lr"],
        "weight_decay": c["weight_decay"], "batch": c["batch"], "epochs": c["epochs"],
        "hidden": "-".join(str(h) for h in c["hidden"]), "dropout": c["dropout"],
        "clip_norm": "none" if c["clip_norm"] is None else c["clip_norm"],
        "precision": c["precision"], "init": c["init"], "seed": c["seed"],
        "step0_loss": s["step0_loss"], "best_val_loss": s["best_val_loss"], "best_epoch": s["best_epoch"],
        "final_train_loss": s["final_train_loss"], "final_val_loss": s["final_val_loss"],
        "val_acc": s["val_acc"], "val_macro_f1": s["val_macro_f1"],
        "time_per_epoch_s": None if s["time_per_epoch_s"] is None else round(s["time_per_epoch_s"], 3),
        "peak_mem_MB": None if s["peak_mem_MB"] is None else round(s["peak_mem_MB"], 1),
        "diverged": "Y" if s["diverged"] else "N",
        "eval_acc": eval_scores["accuracy"] if eval_scores else None,
        "eval_macro_f1": eval_scores["macro_f1"] if eval_scores else None,
        "figure_file": f"figures/{c['exp_id']}.png",
        "notes": "; ".join(p for p in note_parts if p),
    }


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.

    seed_ids     : exp_id các lần chạy baseline khác seed -> cột A của sheet Seeds (tối đa 5 ô A2:A6).
    summary_notes: {group: nhận xét} -> cột "nhận xét ngắn" của sheet Summary.
    """
    import openpyxl

    if len(rows) > MAX_ROWS:
        raise ValueError(f"{len(rows)} dòng vượt quá {MAX_ROWS} dòng có công thức của mẫu")
    ids = [r["exp_id"] for r in rows]
    if len(set(ids)) != len(ids):
        raise ValueError("exp_id bị trùng")

    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]
    header = {cell.value: cell.column for cell in ws[1] if cell.value}
    missing = [k for k in COLUMNS if k not in header]
    if missing:
        raise ValueError(f"mẫu thiếu cột {missing}")
    for i, row in enumerate(rows, start=2):
        for key in COLUMNS:                      # chỉ cột nhập liệu; cột công thức giữ nguyên
            ws.cell(row=i, column=header[key], value=row.get(key))
    for i in range(len(rows) + 2, MAX_ROWS + 2):  # xoá giá trị mẫu còn sót (dòng ví dụ) ở các dòng không dùng
        for key in COLUMNS:
            ws.cell(row=i, column=header[key], value=None)

    if seed_ids is not None:
        seeds = wb["Seeds"]
        if len(seed_ids) > 5:
            raise ValueError("sheet Seeds chỉ có 5 ô A2:A6")
        for i in range(5):
            seeds.cell(row=2 + i, column=1, value=seed_ids[i] if i < len(seed_ids) else None)

    if summary_notes:
        summ = wb["Summary"]
        col = {cell.value: cell.column for cell in summ[1] if cell.value}["nhận xét ngắn (bạn viết)"]
        for r in range(2, summ.max_row + 1):
            g = summ.cell(row=r, column=1).value
            if g in summary_notes:
                summ.cell(row=r, column=col, value=summary_notes[g])

    wb.save(out_path)
